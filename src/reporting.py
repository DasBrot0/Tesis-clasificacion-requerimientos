"""Reports built from finished runs: confusion matrices, comparison tables, Scott-Knott ESD, the
analyses of src/analysis.py and the Excel workbook.

Nothing here trains or imports torch: every file is rebuilt from what a run already saved
(summary.json and pooled_outer_test_predictions.csv), so `scripts/build_report.py` can refresh the
reports of finished runs. Machine-readable files go to outputs/output_raw/<condition>/ and the
readable ones (Excel, figures) to outputs/reports/<condition>/.
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import analysis
from .config import (DATASET_NAMES, DATASETS, LABEL_COLS, MODEL_DISPLAY, MODELS, OUTPUT_RAW, REPORTS,
                     SUBCATEGORY_COLS, run_dir)
from .significance import scott_knott_esd

CONFUSION_CSV = "pooled_confusion_by_label.csv"
SCOTT_KNOTT_CSV = "scott_knott_esd.csv"

# Spanish column titles for the workbook; the order is the order of the columns.
METRIC_TITLES = {
    "macro_f1": "Macro-F1", "weighted_f1": "Weighted-F1", "micro_f1": "Micro-F1",
    "macro_precision": "Macro-precisión", "macro_recall": "Macro-recall",
    "security_recall": "Recall Security", "subset_accuracy": "Exactitud de subconjunto",
    "hamming_loss": "Hamming loss",
}


# --------------------------------------------------------------------------- confusion matrices

def confusion_by_label(pooled: pd.DataFrame) -> pd.DataFrame:
    """One 2x2 matrix per label (TP, FP, FN, TN) from the pooled outer test predictions."""
    rows = []
    for label in LABEL_COLS:
        true = pooled[f"{label}_true"].astype(int)
        pred = pooled[f"{label}_pred"].astype(int)
        rows.append({
            "label": label,
            "TP": int(((true == 1) & (pred == 1)).sum()), "FP": int(((true == 0) & (pred == 1)).sum()),
            "FN": int(((true == 1) & (pred == 0)).sum()), "TN": int(((true == 0) & (pred == 0)).sum()),
        })
    return pd.DataFrame(rows)


def save_confusion_figure(confusion: pd.DataFrame, path: Path, title: str) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, axes = plt.subplots(1, len(confusion), figsize=(3.2 * len(confusion), 3.4))
    for ax, row in zip(axes, confusion.itertuples()):
        matrix = [[row.TN, row.FP], [row.FN, row.TP]]
        ax.imshow(matrix, cmap="Blues")
        peak = max(max(r) for r in matrix)
        for i in range(2):
            for j in range(2):
                ax.text(j, i, matrix[i][j], ha="center", va="center", fontsize=12,
                        color="white" if matrix[i][j] > peak / 2 else "black")
        ax.set(title=row.label, xticks=[0, 1], yticks=[0, 1], xticklabels=["0", "1"],
               yticklabels=["0", "1"], xlabel="predicho", ylabel="real")
    fig.suptitle(title)
    fig.tight_layout()
    fig.savefig(path, dpi=130)
    plt.close(fig)


def save_run_confusion(condition: str, dataset: str, model_name: str) -> pd.DataFrame:
    """CSV next to the run's raw files; figure in outputs/reports/<condition>/confusion_matrices/."""
    out_dir = run_dir(condition, dataset, model_name)
    pooled = pd.read_csv(out_dir / "pooled_outer_test_predictions.csv", encoding="utf-8-sig", dtype={"ID": str})
    confusion = confusion_by_label(pooled)
    confusion.to_csv(out_dir / CONFUSION_CSV, index=False, encoding="utf-8-sig")
    figures = REPORTS / condition / "confusion_matrices"
    figures.mkdir(parents=True, exist_ok=True)
    save_confusion_figure(confusion, figures / f"{dataset}_{model_name}.png",
                          f"{DATASET_NAMES[dataset]} — {MODEL_DISPLAY[model_name]} ({condition})")
    return confusion


# --------------------------------------------------------------------------- finished runs

def finished_runs(condition: str) -> list[tuple[str, str, dict]]:
    """(dataset, model, summary) of every finished run of a condition, in the canonical order."""
    runs = []
    for dataset in DATASETS:
        for model_name in MODELS:
            path = run_dir(condition, dataset, model_name) / "summary.json"
            if path.exists():
                runs.append((dataset, model_name, json.loads(path.read_text(encoding="utf-8"))))
    return runs


def run_confusion(condition: str, dataset: str, model_name: str) -> pd.DataFrame:
    path = run_dir(condition, dataset, model_name) / CONFUSION_CSV
    if path.exists():
        return pd.read_csv(path, encoding="utf-8-sig")
    return save_run_confusion(condition, dataset, model_name)


def update_scott_knott(condition: str) -> pd.DataFrame:
    """Scott-Knott ESD of the models of each dataset over their outer-fold macro-F1."""
    parts = []
    runs = finished_runs(condition)
    for dataset in DATASETS:
        samples = {MODEL_DISPLAY[m]: [fold["test"]["macro_f1"] for fold in report["outer_folds"]]
                   for d, m, report in runs if d == dataset}
        if len(samples) < 2:
            continue
        ranking = scott_knott_esd(samples)
        ranking.insert(0, "dataset", DATASET_NAMES[dataset])
        parts.append(ranking.rename(columns={"treatment": "model"}))
    table = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
    path = OUTPUT_RAW / condition / SCOTT_KNOTT_CSV
    path.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(path, index=False, encoding="utf-8-sig")
    return table


def update_comparison_table(condition: str) -> pd.DataFrame:
    rows = []
    for dataset, model_name, report in finished_runs(condition):
        outer, pooled = report["outer_aggregate"], report["pooled_outer_test"]
        fmt = lambda name: f"{outer[name]['mean']:.4f} ± {outer[name]['std']:.4f}"
        rows.append({
            "Condición": condition, "Dataset": DATASET_NAMES[dataset], "Modelo": MODEL_DISPLAY[model_name],
            "Macro-F1 (5 folds externos)": fmt("macro_f1"), "Weighted-F1 (5 folds externos)": fmt("weighted_f1"),
            "Recall Security (5 folds externos)": fmt("security_recall"),
            "Macro-F1 agrupado": pooled["macro_f1"], "Weighted-F1 agrupado": pooled["weighted_f1"],
            "Micro-F1 agrupado": pooled["micro_f1"], "Recall Security agrupado": pooled["security_recall"],
            **{f"F1 {row['label']}": row["f1"] for row in report["pooled_outer_test_per_label"]},
            "Etiquetas en el promedio": ", ".join(report["labels_in_macro_average"]),
            "Sintéticos": report["synthetic_rows"], "Minutos": report["elapsed_minutes"],
        })
    table = pd.DataFrame(rows)
    path = OUTPUT_RAW / condition / "model_comparison.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(path, index=False, encoding="utf-8-sig")
    return table


# --------------------------------------------------------------------------- Excel workbook

header_fill = PatternFill("solid", fgColor="1F4E78")
group_fill = PatternFill("solid", fgColor="D9E1F2")
hit_fill = PatternFill("solid", fgColor="E2EFDA")
miss_fill = PatternFill("solid", fgColor="FCE4D6")
header_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
score_scale = dict(start_type="num", start_value=0, start_color="F8696B",
                   mid_type="num", mid_value=0.5, mid_color="FFEB84",
                   end_type="num", end_value=1, end_color="63BE7B")


def write_table(ws, top: int, headers: list[str], rows: list[list], formats: dict[int, str]) -> int:
    """Header + rows starting at `top`; returns the last row written."""
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(top, col, title)
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border
    for r, values in enumerate(rows, start=top + 1):
        for col, value in enumerate(values, start=1):
            if isinstance(value, float) and np.isnan(value):
                value = None
            cell = ws.cell(r, col, value)
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if isinstance(value, str) else "center")
            if col in formats:
                cell.number_format = formats[col]
    ws.row_dimensions[top].height = 32
    return top + len(rows)


def bold_best(ws, first: int, last: int, col: int, key_col: int = 1, lower_is_better: bool = False) -> None:
    """Bold the best value of `col` within each block of rows that share the value of `key_col`."""
    blocks: dict[str, list[int]] = {}
    for r in range(first, last + 1):
        blocks.setdefault(ws.cell(r, key_col).value, []).append(r)
    for rows in blocks.values():
        pick = min if lower_is_better else max
        best = pick(rows, key=lambda r: ws.cell(r, col).value)
        ws.cell(best, col).font = Font(bold=True)


def widths(ws, values: dict[int, float]) -> None:
    for col, width in values.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def sheet_title(ws, title: str, note: str = "", note_columns: str = "A2:K2", note_height: int = 32) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(note_columns)
        ws.row_dimensions[2].height = note_height


def export_excel(condition: str) -> Path | None:
    """outputs/reports/<condition>/results_<condition>.xlsx: general metrics, metrics by class,
    confusion matrices, statistical comparisons and analyses of every finished run, formatted to
    be read as is."""
    runs = finished_runs(condition)
    if not runs:
        return None

    wb = Workbook()

    # 1) General metrics: pooled outer test + mean ± std of the five outer folds.
    ws = wb.active
    ws.title = "Resumen general"
    ws["A1"] = f"Métricas generales — condición {condition}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Agrupado: predicciones de prueba de los 5 folds externos juntas (cada requisito real "
                "una vez). Folds: media ± desviación estándar de los 5 folds externos. "
                "En negrita, el mejor modelo de cada dataset.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:P2")
    ws.row_dimensions[2].height = 32
    names = list(METRIC_TITLES)
    ws.cell(4, 3, "Agrupado (5 folds externos)")
    ws.merge_cells(start_row=4, start_column=3, end_row=4, end_column=2 + len(names))
    ws.cell(4, 3 + len(names), "Media ± desv. de los 5 folds externos")
    ws.merge_cells(start_row=4, start_column=3 + len(names), end_row=4, end_column=2 + 2 * len(names))
    for col in (3, 3 + len(names)):
        ws.cell(4, col).fill, ws.cell(4, col).font, ws.cell(4, col).alignment = group_fill, Font(bold=True), center
    headers = ["Dataset", "Modelo", *METRIC_TITLES.values(), *METRIC_TITLES.values()]
    rows = []
    for dataset, model_name, report in runs:
        pooled, outer = report["pooled_outer_test"], report["outer_aggregate"]
        rows.append([DATASET_NAMES[dataset], MODEL_DISPLAY[model_name], *(pooled[n] for n in names),
                     *(f"{outer[n]['mean']:.3f} ± {outer[n]['std']:.3f}" for n in names)])
    last = write_table(ws, 5, headers, rows, {c: "0.000" for c in range(3, 3 + len(names))})
    for offset, name in enumerate(names):
        bold_best(ws, 6, last, 3 + offset, lower_is_better=name == "hamming_loss")
    # Hamming loss (last pooled column) is better when lower: it stays out of the color scale.
    ws.conditional_formatting.add(f"C6:{get_column_letter(2 + len(names) - 1)}{last}", ColorScaleRule(**score_scale))
    widths(ws, {1: 14, 2: 14, **{c: 13 for c in range(3, 3 + 2 * len(names))}})
    ws.freeze_panes = "C6"

    # 2) F1 by class, one row per (dataset, model): the compact view to compare models.
    ws = wb.create_sheet("F1 por clase")
    ws["A1"] = "F1 por clase (predicciones externas agrupadas)"
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Dataset", "Modelo", *LABEL_COLS, "Macro-F1"]
    rows = []
    for dataset, model_name, report in runs:
        f1 = {row["label"]: row["f1"] for row in report["pooled_outer_test_per_label"]}
        rows.append([DATASET_NAMES[dataset], MODEL_DISPLAY[model_name], *(f1[l] for l in LABEL_COLS),
                     report["pooled_outer_test"]["macro_f1"]])
    last = write_table(ws, 3, headers, rows, {c: "0.000" for c in range(3, 4 + len(LABEL_COLS))})
    for col in range(3, 4 + len(LABEL_COLS)):
        bold_best(ws, 4, last, col)
    ws.conditional_formatting.add(f"C4:{get_column_letter(3 + len(LABEL_COLS))}{last}", ColorScaleRule(**score_scale))
    widths(ws, {1: 14, 2: 14, **{c: 16 for c in range(3, 4 + len(LABEL_COLS))}})
    ws.freeze_panes = "C4"

    # 3) Every metric by class, with the confusion counts behind it.
    ws = wb.create_sheet("Métricas por clase")
    ws["A1"] = "Métricas por clase (predicciones externas agrupadas)"
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Dataset", "Modelo", "Clase", "Precisión", "Recall", "F1", "Soporte",
               "VP", "FP", "FN", "VN", "En el promedio"]
    rows = []
    for dataset, model_name, report in runs:
        confusion = run_confusion(condition, dataset, model_name).set_index("label")
        for row in report["pooled_outer_test_per_label"]:
            counts = confusion.loc[row["label"]]
            rows.append([DATASET_NAMES[dataset], MODEL_DISPLAY[model_name], row["label"], row["precision"],
                         row["recall"], row["f1"], int(row["support"]), int(counts["TP"]), int(counts["FP"]),
                         int(counts["FN"]), int(counts["TN"]), "Sí" if row["in_macro_average"] else "No"])
    last = write_table(ws, 3, headers, rows, {4: "0.000", 5: "0.000", 6: "0.000"})
    ws.conditional_formatting.add(f"D4:F{last}", ColorScaleRule(**score_scale))
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{last}"
    widths(ws, {1: 14, 2: 14, 3: 17, **{c: 11 for c in range(4, len(headers) + 1)}, len(headers): 15})
    ws.freeze_panes = "D4"

    # 4) Confusion matrices: one block per run, one 2x2 matrix per class.
    ws = wb.create_sheet("Matrices de confusión")
    ws["A1"] = "Matrices de confusión por clase (predicciones externas agrupadas)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Filas: valor real; columnas: valor predicho. Verde: aciertos (VN, VP); naranja: errores (FP, FN)."
    top = 4
    for dataset, model_name, _ in runs:
        confusion = run_confusion(condition, dataset, model_name)
        ws.cell(top, 1, f"{DATASET_NAMES[dataset]} — {MODEL_DISPLAY[model_name]}").font = Font(bold=True, size=12)
        for k, row in enumerate(confusion.itertuples()):
            left = 1 + 4 * k
            ws.cell(top + 1, left, row.label)
            ws.merge_cells(start_row=top + 1, start_column=left, end_row=top + 1, end_column=left + 2)
            title = ws.cell(top + 1, left)
            title.fill, title.font, title.alignment = header_fill, header_font, center
            cells = [["Real \\ Pred.", "0", "1"], ["0", row.TN, row.FP], ["1", row.FN, row.TP]]
            for i, values in enumerate(cells):
                for j, value in enumerate(values):
                    cell = ws.cell(top + 2 + i, left + j, value)
                    cell.border, cell.alignment = border, center
                    if i == 0 or j == 0:
                        cell.fill, cell.font = group_fill, Font(bold=True)
                    else:
                        cell.fill = hit_fill if i == j else miss_fill
        top += 6
    widths(ws, {c: (12 if c % 4 == 1 else 9 if c % 4 else 3) for c in range(1, 4 * len(LABEL_COLS) + 1)})

    # 5) Each outer fold, for the variability behind the mean ± std.
    ws = wb.create_sheet("Folds externos")
    ws["A1"] = "Métricas de cada fold externo"
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Dataset", "Modelo", "Fold externo", *METRIC_TITLES.values()]
    rows = [[DATASET_NAMES[dataset], MODEL_DISPLAY[model_name], fold["outer"] + 1,
             *(fold["test"][n] for n in names)]
            for dataset, model_name, report in runs for fold in report["outer_folds"]]
    last = write_table(ws, 3, headers, rows, {c: "0.000" for c in range(4, 4 + len(names))})
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{last}"
    widths(ws, {1: 14, 2: 14, 3: 11, **{c: 13 for c in range(4, 4 + len(names))}})
    ws.freeze_panes = "D4"

    # 6) Scott-Knott ESD of the models of each dataset (outer-fold macro-F1).
    ranking = update_scott_knott(condition)
    if not ranking.empty:
        ws = wb.create_sheet("Scott-Knott ESD")
        ws["A1"] = "Scott-Knott ESD por dataset (macro-F1 de los 5 folds externos)"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = ("Rango 1 = grupo superior. Los modelos con el mismo rango no difieren de forma "
                    "significativa (Kruskal-Wallis, α = 0,05) o su diferencia es despreciable "
                    "(δ de Cliff < 0,147). δ, magnitud y p comparan el grupo con el siguiente.")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("A2:K2")
        ws.row_dimensions[2].height = 32
        headers = ["Dataset", "Modelo", "Rango", "Grupo", "Mediana", "Media", "Desv.", "Folds",
                   "δ de Cliff vs. siguiente", "Magnitud", "p vs. siguiente"]
        rows = [[r.dataset, r.model, r.rank, r.group, r.median, r.mean, r.std, r.n,
                 None if pd.isna(r.delta_vs_next_group) else r.delta_vs_next_group,
                 r.magnitude_vs_next_group, None if pd.isna(r.p_vs_next_group) else r.p_vs_next_group]
                for r in ranking.itertuples()]
        last = write_table(ws, 4, headers, rows, {5: "0.000", 6: "0.000", 7: "0.000", 9: "0.000", 11: "0.0000"})
        for r in range(5, last + 1):
            if ws.cell(r, 3).value == 1:
                for col in (2, 3):
                    ws.cell(r, col).font = Font(bold=True)
                    ws.cell(r, col).fill = hit_fill
        widths(ws, {1: 14, 2: 14, 3: 8, 4: 36, 5: 10, 6: 10, 7: 10, 8: 8, 9: 14, 10: 13, 11: 13})
        ws.freeze_panes = "C5"

    analysis_sheets(wb, condition, runs, ranking)

    path = REPORTS / condition / f"results_{condition}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def analysis_sheets(wb, condition: str, runs: list[tuple[str, str, dict]], ranking: pd.DataFrame) -> None:
    """Sheets of src/analysis.py: model tests, subcategory confusion, misclassified requirements,
    best model, Transformers vs Deep Learning and (C1) the diversity of the synthetics."""
    significant_fill = PatternFill("solid", fgColor="C6EFCE")

    # 7) Cochran's Q and pairwise McNemar between the models of each dataset.
    omnibus, pairs = analysis.model_tests(condition)
    if not omnibus.empty:
        ws = wb.create_sheet("Cochran Q y McNemar")
        sheet_title(ws, "Diferencias entre modelos por clase (predicciones externas agrupadas)",
                    "Cochran Q: ¿los modelos del dataset difieren en sus aciertos en esa clase? Solo si "
                    "p < 0,05 se interpretan las comparaciones por pares (McNemar exacto, p ajustado por "
                    "Holm sobre los pares de esa clase y dataset). En verde, las diferencias significativas.",
                    "A2:M2", 45)
        accuracy_cols = [c for c in omnibus.columns if c.startswith("accuracy ")]
        headers = ["Dataset", "Clase", "En el promedio", "Soporte", "Q de Cochran", "p", "Significativa",
                   *(f"Exactitud {c.removeprefix('accuracy ')}" for c in accuracy_cols)]
        rows = [[r["dataset"], r["label"], "Sí" if r["in_macro_average"] else "No", r["support"],
                 r["cochran_q"], r["p"], "Sí" if r["significant"] else "No", *(r[c] for c in accuracy_cols)]
                for r in omnibus.to_dict("records")]
        ws.cell(4, 1, "Prueba ómnibus").font = Font(bold=True, size=12)
        last = write_table(ws, 5, headers, rows,
                           {5: "0.00", 6: "0.0000", **{c: "0.000" for c in range(8, 8 + len(accuracy_cols))}})
        for r in range(6, last + 1):
            if ws.cell(r, 7).value == "Sí":
                for col in (5, 6, 7):
                    ws.cell(r, col).fill = significant_fill
        top = last + 3
        ws.cell(top - 1, 1, "Comparaciones por pares (McNemar exacto con Holm)").font = Font(bold=True, size=12)
        headers = ["Dataset", "Clase", "Modelo A", "Modelo B", "F1 A", "F1 B", "Solo A acierta",
                   "Solo B acierta", "p", "p Holm", "Significativa", "Más aciertos"]
        rows = [[r["dataset"], r["label"], r["model_a"], r["model_b"], r["f1_a"], r["f1_b"], r["only_a_right"],
                 r["only_b_right"], r["p"], r["p_holm"], "Sí" if r["significant"] else "No", r["more_hits"]]
                for r in pairs.to_dict("records")]
        last = write_table(ws, top, headers, rows, {5: "0.000", 6: "0.000", 9: "0.0000", 10: "0.0000"})
        for r in range(top + 1, last + 1):
            if ws.cell(r, 11).value == "Sí":
                for col in (10, 11, 12):
                    ws.cell(r, col).fill = significant_fill
        ws.auto_filter.ref = f"A{top}:L{last}"
        widths(ws, {1: 14, 2: 16, 3: 14, 4: 14, **{c: 12 for c in range(5, 14)}})

    # 8) Co-confusion between subcategories, one matrix per run.
    confusion = analysis.confusion_tables(condition)
    if not confusion.empty:
        ws = wb.create_sheet("Confusión subcategorías")
        sheet_title(ws, "Confusiones entre subcategorías (predicciones externas agrupadas)",
                    "Filas: subcategoría real; columnas: subcategoría predicha. Diagonal (verde): acierto. "
                    "Fuera de la diagonal (naranja): el modelo omitió la subcategoría de la fila y predijo la "
                    "de la columna, que el requisito no tiene. Columna «Ninguna»: la omitió sin predecir otra. "
                    "Fila «Ninguna»: requisitos sin subcategoría real; cuenta las subcategorías predichas de "
                    "más, y en «Ninguna» los que quedaron correctamente sin subcategoría.",
                    "A2:AB2", 48)
        names = [*SUBCATEGORY_COLS, analysis.NO_SUBCATEGORY]
        span = len(names) + 2
        top = 4
        for dataset in confusion["dataset"].unique():
            part = confusion[confusion["dataset"] == dataset]
            for k, model in enumerate(part["model"].unique()):
                matrix = part[part["model"] == model].pivot(index="real", columns="predicted", values="count")
                left = 1 + k * span
                ws.cell(top, left, f"{dataset} — {model}").font = Font(bold=True, size=12)
                corner = ws.cell(top + 1, left, "Real \\ Predicha")
                corner.fill, corner.font, corner.alignment, corner.border = group_fill, Font(bold=True), center, border
                for j, name in enumerate(names):
                    for cell in (ws.cell(top + 1, left + 1 + j, name), ws.cell(top + 2 + j, left, name)):
                        cell.fill, cell.font, cell.alignment, cell.border = group_fill, Font(bold=True), center, border
                for i, real in enumerate(names):
                    for j, predicted in enumerate(names):
                        value = int(matrix.loc[real, predicted])
                        cell = ws.cell(top + 2 + i, left + 1 + j, value)
                        cell.border, cell.alignment = border, center
                        if i == j:
                            cell.fill = hit_fill
                        elif value and predicted != analysis.NO_SUBCATEGORY:
                            cell.fill = miss_fill
            top += len(names) + 4
        widths(ws, {c: (16 if c % span == 1 else 3 if c % span == 0 else 11) for c in range(1, 4 * span + 1)})

    # 9) Misclassified requirements, for the manual error review.
    errors = analysis.misclassified(condition)
    if not errors.empty:
        ws = wb.create_sheet("Errores para revisión")
        sheet_title(ws, "Requisitos con al menos un error (predicciones externas agrupadas)",
                    "Una fila por requisito real que algún modelo clasificó mal en alguna clase. Ordenado por "
                    "el número de modelos que fallan: los primeros son los más difíciles o los que podrían "
                    "estar mal etiquetados.", "A2:I2")
        model_cols = [c for c in errors.columns if c not in ("dataset", "ID", "models_wrong", "real", "requirement")]
        headers = ["Dataset", "ID", "Modelos que fallan", "Etiquetas reales", *model_cols, "Requisito"]
        rows = [[r["dataset"], r["ID"], r["models_wrong"], r["real"], *(r[c] for c in model_cols), r["requirement"]]
                for r in errors.to_dict("records")]
        last = write_table(ws, 4, headers, rows, {})
        for r in range(5, last + 1):
            truth = ws.cell(r, 4).value
            for col in range(5, 5 + len(model_cols)):
                ws.cell(r, col).fill = hit_fill if ws.cell(r, col).value == truth else miss_fill
            ws.cell(r, len(headers)).alignment = Alignment(wrap_text=True, vertical="top")
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
        widths(ws, {1: 13, 2: 15, 3: 10, 4: 28, **{c: 28 for c in range(5, 5 + len(model_cols))},
                    len(headers): 90})
        ws.freeze_panes = "E5"

    # 10) Best model of each dataset.
    best = analysis.best_models(condition, ranking)
    if not best.empty:
        ws = wb.create_sheet("Mejor modelo")
        sheet_title(ws, "Mejor modelo por dataset",
                    "Regla (METODOLOGIA §4.3.2): los modelos del grupo superior de Scott-Knott ESD no "
                    "difieren en macro-F1; entre ellos se elige el de mayor recall de Security (agrupado), "
                    "porque omitir un requisito de seguridad es el error más costoso.",
                    "A2:G2", 32)
        headers = ["Dataset", "Mejor modelo", "Macro-F1", "Recall Security", "Motivo",
                   "Grupo superior SK-ESD", "Candidatos (por recall de Security)"]
        rows = [[r["dataset"], r["best_model"], r["macro_f1"], r["security_recall"], r["reason"], r["top_group"],
                 r["candidates"]] for r in best.to_dict("records")]
        last = write_table(ws, 4, headers, rows, {3: "0.000", 4: "0.000"})
        for r in range(5, last + 1):
            ws.cell(r, 2).font, ws.cell(r, 2).fill = Font(bold=True), hit_fill
            for col in (5, 7):
                ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
        widths(ws, {1: 14, 2: 15, 3: 10, 4: 11, 5: 45, 6: 30, 7: 70})

    # 11) Transformers against Deep Learning.
    by_dataset, by_label = analysis.family_comparison(condition, runs, ranking)
    if not by_dataset.empty:
        ws = wb.create_sheet("Transformers vs DL")
        sheet_title(ws, "Transformers (BETO, RoBERTa-BNE) frente a Deep Learning (CNN, BiLSTM)",
                    "Macro-F1 de cada familia: media de sus dos modelos (agrupado). δ de Cliff y Mann-Whitney "
                    "comparan los 10 macro-F1 de folds externos de cada familia (2 modelos × 5 folds); δ > 0 "
                    "favorece a los Transformers. Rango SK: el mejor rango Scott-Knott ESD de la familia.",
                    "A2:L2", 45)
        headers = ["Dataset", "Macro-F1 Transformers", "Macro-F1 Deep Learning", "Δ macro-F1",
                   "Mejor Transformer", "Mejor DL", "δ de Cliff (folds)", "Magnitud", "p Mann-Whitney",
                   "Rango SK Transformers", "Rango SK DL"]
        keys = ["dataset", "macro_f1 Transformers (media agrupada)", "macro_f1 Deep Learning (media agrupada)",
                "delta_macro_f1", "mejor Transformers", "mejor Deep Learning", "cliffs_delta_folds", "magnitude",
                "mannwhitney_p", "mejor rango SK Transformers", "mejor rango SK Deep Learning"]
        rows = [[r.get(k) for k in keys] for r in by_dataset.to_dict("records")]
        last = write_table(ws, 4, headers, rows, {2: "0.000", 3: "0.000", 4: "+0.000;-0.000", 7: "+0.000;-0.000",
                                                   9: "0.0000", 10: "0", 11: "0"})
        for col in range(1, len(headers) + 1):
            ws.cell(last, col).font = Font(bold=True)
        top = last + 3
        ws.cell(top - 1, 1, "F1 por clase (media de los dos modelos de cada familia)").font = Font(bold=True, size=12)
        headers = ["Dataset", "Clase", "Soporte", "F1 Transformers", "F1 Deep Learning", "Δ F1"]
        rows = [[r["dataset"], r["label"], r["support"], r["F1 Transformers"], r["F1 Deep Learning"], r["delta_f1"]]
                for r in by_label.to_dict("records")]
        last = write_table(ws, top, headers, rows, {3: "0", 4: "0.000", 5: "0.000", 6: "+0.000;-0.000"})
        ws.conditional_formatting.add(f"D{top + 1}:E{last}", ColorScaleRule(**score_scale))
        widths(ws, {1: 20, 2: 16, **{c: 14 for c in range(3, 12)}, 8: 28})

    # 12) Diversity of the synthetics (C1 only).
    if condition == "c1_augmented":
        diversity = analysis.synthetic_diversity()
        if not diversity.empty:
            ws = wb.create_sheet("Diversidad sintéticos")
            sheet_title(ws, "Diversidad de los sintéticos aceptados (similitud coseno LaBSE)",
                        "Con su fuente y con el real más parecido: calculadas por el filtro G5 (umbrales: casi "
                        "copia ≥ 0,85, fuga ≥ 0,90). Pares sintético-sintético, real-real y sintético-real: "
                        "media sobre todos los pares, de scripts/augmentation/g9_diversity.py; si el valor "
                        "sintético-sintético supera mucho al real-real, las variantes son repetitivas.",
                        "A2:L2", 48)
            titles = {
                "dataset": "Dataset", "synthetics": "Sintéticos", "sources": "Fuentes",
                "variants_per_source": "Variantes por fuente", "sim_source_mean": "Sim. con su fuente (media)",
                "sim_source_max": "Sim. con su fuente (máx.)",
                "sim_max_real_mean": "Sim. con el real más parecido (media)",
                "sim_max_real_max": "Sim. con el real más parecido (máx.)", "real_security": "Reales de seguridad",
                "sim_synthetic_synthetic": "Sintético-sintético (media)", "sim_real_real": "Real-real (media)",
                "sim_synthetic_real": "Sintético-real (media)",
            }
            columns = [c for c in titles if c in diversity.columns]
            rows = [[r[c] for c in columns] for r in diversity.to_dict("records")]
            last = write_table(ws, 4, [titles[c] for c in columns], rows,
                               {k + 1: "0.000" for k, c in enumerate(columns) if c.startswith(("sim", "variants"))})
            if "sim_synthetic_synthetic" not in columns:
                ws.cell(last + 2, 1, "Falta correr scripts/augmentation/g9_diversity.py para las similitudes "
                                     "entre pares (necesita transformers y LaBSE).").font = Font(italic=True)
            widths(ws, {k + 1: (14 if k == 0 else 16) for k in range(len(columns))})
